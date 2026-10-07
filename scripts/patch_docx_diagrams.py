#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script cập nhật trực tiếp (in-place) file DOCX hiện hành:
- Đọc file DOCX đang có (giữ nguyên 100% mọi chỉnh sửa người dùng đã gõ trong WPS Office/Word).
- Tìm đúng 2 khối bảng chứa mã raw Mermaid (graph TD và sequenceDiagram).
- Thay thế 2 khối đó bằng ảnh sơ đồ thực tế chất lượng cao (High-DPI).
- Lưu đè lại chính file đó mà KHÔNG tạo lại từ đầu.
"""

import sys
import time
from pathlib import Path

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def patch_docx():
    base_dir = Path(__file__).resolve().parent.parent
    docx_path = base_dir / "docs" / "product" / "PRODUCT_REQUIREMENTS_DOCUMENT_LINGUAL.docx"
    assets_dir = base_dir / "docs" / "assets"

    img_pillars = assets_dir / "diagram_pillars.png"
    img_sequence = assets_dir / "diagram_sequence.png"

    if not docx_path.exists():
        print(f"❌ Không tìm thấy file: {docx_path}")
        return False

    print(f"📄 Đang nạp tài liệu hiện tại: {docx_path}")
    
    # Kiểm tra xem file có bị khóa bởi ứng dụng ngoài (WPS/Word) không
    try:
        with open(docx_path, "r+b"):
            pass
    except PermissionError:
        print("⚠️ CẢNH BÁO: File đang bị khóa bởi WPS Office / Word!")
        print("👉 Vui lòng nhấn Ctrl + S trong WPS Office để LƯU THAY ĐỔI của bạn, sau đó ĐÓNG tab tài liệu hoặc đóng WPS Office để giải phóng file.")
        return False

    doc = docx.Document(str(docx_path))
    replaced_count = 0

    for t in list(doc.tables):
        try:
            cell_text = t.cell(0, 0).text.strip()
        except Exception:
            continue

        if cell_text.startswith("graph TD"):
            if img_pillars.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = docx.shared.Pt(12)
                p.paragraph_format.space_after = docx.shared.Pt(12)
                run = p.add_run()
                run.add_picture(str(img_pillars), width=Inches(5.8))
                
                # Chèn ảnh vào đúng vị trí của bảng cũ
                t._element.addprevious(p._element)
                t._element.getparent().remove(t._element)
                replaced_count += 1
                print("✅ Đã thay thế mã 'graph TD' bằng hình ảnh diagram_pillars.png")

        elif cell_text.startswith("sequenceDiagram"):
            if img_sequence.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = docx.shared.Pt(12)
                p.paragraph_format.space_after = docx.shared.Pt(12)
                run = p.add_run()
                run.add_picture(str(img_sequence), width=Inches(6.2))
                
                # Chèn ảnh vào đúng vị trí của bảng cũ
                t._element.addprevious(p._element)
                t._element.getparent().remove(t._element)
                replaced_count += 1
                print("✅ Đã thay thế mã 'sequenceDiagram' bằng hình ảnh diagram_sequence.png")

    if replaced_count > 0:
        doc.save(str(docx_path))
        print(f"🎉 HOÀN TẤT: Đã cập nhật {replaced_count} sơ đồ trực tiếp vào file. Toàn bộ nội dung chỉnh sửa của anh được giữ nguyên 100%!")
        return True
    else:
        print("ℹ️ Không tìm thấy bảng mã Mermaid thô nào cần thay thế (có thể đã được cập nhật trước đó).")
        return True


if __name__ == "__main__":
    patch_docx()
