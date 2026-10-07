#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script cập nhật trực tiếp (in-place) file DOCX SPRINT_PLAN_LINGUAL_MCS2026.docx:
- Đọc file DOCX đang có (giữ nguyên 100% mọi chỉnh sửa người dùng đã gõ trong WPS Office/Word).
- Tìm đúng bảng chứa khối text ASCII ma trận 3 luồng (TUẦN 1 - 2, LUỒNG 1...).
- Thay thế bảng text bị vỡ đó bằng ảnh sơ đồ Ma trận 3 Luồng sắc nét (diagram_workstreams_matrix.png).
- Lưu đè lại chính file đó mà KHÔNG tạo lại từ đầu.
"""

import sys
from pathlib import Path

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def patch_sprint_docx():
    base_dir = Path(__file__).resolve().parent.parent
    docx_path = base_dir / "docs" / "product" / "SPRINT_PLAN_LINGUAL_MCS2026.docx"
    assets_dir = base_dir / "docs" / "assets"

    img_matrix = assets_dir / "diagram_workstreams_matrix.png"
    img_timeline = assets_dir / "diagram_sprint_timeline.png"

    if not docx_path.exists():
        print(f"❌ Không tìm thấy file: {docx_path}")
        return False

    print(f"📄 Đang nạp tài liệu hiện tại: {docx_path}")

    # Kiểm tra xem file có bị khóa bởi WPS Office / Word không
    try:
        with open(docx_path, "r+b"):
            pass
    except PermissionError:
        print("⚠️ CẢNH BÁO: File đang bị khóa bởi WPS Office / Word!")
        print("👉 Vui lòng nhấn Ctrl + S trong WPS Office để LƯU THAY ĐỔI của bạn, sau đó ĐÓNG tab tài liệu hoặc đóng WPS Office để giải phóng file.")
        return False

    doc = docx.Document(str(docx_path))
    replaced_count = 0

    target_img = img_matrix if img_matrix.exists() else img_timeline

    for t in list(doc.tables):
        try:
            cell_text = t.cell(0, 0).text
        except Exception:
            continue

        # Tìm bảng chứa đoạn text ma trận 3 luồng ASCII
        if "TUẦN 1 - 2" in cell_text and "LUỒNG 1" in cell_text:
            if target_img.exists():
                # Tạo đoạn chứa ảnh sơ đồ
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(14)
                p_img.paragraph_format.space_after = Pt(6)
                run_img = p_img.add_run()
                run_img.add_picture(str(target_img), width=Inches(6.2))

                # Tạo đoạn chú thích ảnh (Caption)
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(0)
                p_cap.paragraph_format.space_after = Pt(16)
                run_cap = p_cap.add_run("Hình: Sơ đồ Ma trận 3 Luồng Phát triển Song song & Lộ trình 10 Tuần của LINGUAL (MCS 2026)")
                run_cap.italic = True
                run_cap.font.name = "Calibri"
                run_cap.font.size = Pt(9)
                run_cap.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

                # Chèn ảnh và caption vào ngay vị trí của bảng cũ
                t._element.addprevious(p_img._element)
                t._element.addprevious(p_cap._element)
                t._element.getparent().remove(t._element)
                replaced_count += 1
                print("✅ Đã thay thế bảng text ASCII bị vỡ bằng hình ảnh sơ đồ Ma trận 3 Luồng!")

    if replaced_count > 0:
        doc.save(str(docx_path))
        print(f"🎉 HOÀN TẤT: Đã cập nhật sơ đồ trực tiếp vào file. Toàn bộ nội dung chỉnh sửa của bạn được giữ nguyên 100%!")
        return True
    else:
        print("ℹ️ Không tìm thấy bảng text ASCII nào cần thay thế (có thể đã được cập nhật trước đó).")
        return True


if __name__ == "__main__":
    patch_sprint_docx()
